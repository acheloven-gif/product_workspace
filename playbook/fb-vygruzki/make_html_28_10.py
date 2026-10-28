# HTML-подборки по итоговой таблице FB-выгрузок (28.10).
# Вход: reports/fb-kandidaty-itog-28-10.xlsx → выход: reports/fb-shortlist-28-10.html
# Без JS (вкладки на CSS) — открывается из вложения ClickUp.
import openpyxl, sys, html

ROOT = sys.argv[1] if len(sys.argv) > 1 else '.'
wb = openpyxl.load_workbook(f'{ROOT}/reports/fb-kandidaty-itog-28-10.xlsx', data_only=True)

def rows(name):
    ws = wb[name]; hdr = [c.value for c in ws[3]]
    out = []
    for r in range(4, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, len(hdr) + 1)]
        if v[1]: out.append(dict(zip(hdr, v)))
    return out

e = lambda s: html.escape(str(s)) if s is not None else ''
def money(v):
    if not isinstance(v, (int, float)): return '—'
    return f'${v/1e6:.2f}M' if v >= 1e6 else f'${v/1e3:.0f}k'
def zone(p):
    if not isinstance(p, (int, float)): return 'na', 'н/д'
    return ('lo', 'дёшево') if p <= 1/3 else (('mid', 'средне') if p <= 2/3 else ('hi', 'дорого'))
def links(d):
    out = []
    for k, lbl in [('Воронка', 'воронка'), ('App Store', 'App Store'), ('Google Play', 'Google Play'), ('Топ-объявление (Ad Library)', 'объявление')]:
        u = d.get(k)
        if isinstance(u, str) and u.startswith('http'): out.append(f'<a href="{e(u)}" target="_blank" rel="noopener">{lbl}</a>')
    return ' · '.join(out)

def table(items, why_key):
    tr = []
    for d in items:
        zc, zl = zone(d.get('CPA-перцентиль окт (0 = дешевле всех)'))
        cpa = d.get('CPA окт')
        cpa_s = f'${cpa:.0f}' if isinstance(cpa, (int, float)) else 'н/д'
        if not isinstance(cpa, (int, float)): zc, zl = 'na', 'оплат нет'
        t = d.get('Тренд авг→окт') or '—'
        sc = d.get('Скор (0–100)')
        tr.append(f'''<tr>
<td class="n">{e(d['#'])}</td>
<td class="app"><b>{e(d['Приложение'])}</b><span>{e(d.get('Про что приложение'))}</span><span class="lk">{links(d)}</span></td>
<td class="num"><b>{f"{sc:.0f}" if isinstance(sc, (int, float)) else "—"}</b></td>
<td class="num">{money(d.get('Спенд 28д, Октябрь (10.10)'))}</td>
<td class="num"><i class="z {zc}"></i>{cpa_s}<small>{zl}</small></td>
<td><em class="t t-{e(t)}">{e(t)}</em></td>
<td>{e(d.get('Тир гео'))}<small>{e(d.get('Гео (доли стран)'))}</small></td>
<td class="why">{e(d.get(why_key))}</td>
</tr>''')
    return f'''<div class="wrap"><table>
<thead><tr><th>#</th><th>Приложение</th><th>Скор</th><th>Спенд 28д, окт</th><th>CPA окт</th><th>Тренд авг→окт</th><th>Гео</th><th>{'Почему под вопросом' if why_key != 'Почему' else 'Почему'}</th></tr></thead>
<tbody>{''.join(tr)}</tbody></table></div>'''

S = rows('Шорт-лист'); Q = rows('Под вопросом'); T = rows('Топ-кандидаты')

page = f'''<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Шорт-лист FB-выгрузок</title>
<style>
:root{{--bg:#fff8f3;--card:#ffffff;--ink:#3a2a24;--mute:#8a7469;--line:#f1dfd4;--accent:#e07a5f;--accent-soft:#fde5da;
--lo:#5aa469;--mid:#e0a43a;--hi:#d1495b;--na:#b8a89f}}
@media (prefers-color-scheme: dark){{:root:not([data-theme="light"]){{--bg:#1f1916;--card:#2a221e;--ink:#f3e7e0;--mute:#b9a499;--line:#3d322c;--accent:#f0957a;--accent-soft:#3a2a24}}}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--bg);color:var(--ink);font:15px/1.45 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif}}
main{{max-width:1240px;margin:0 auto;padding:28px 16px 48px}}
h1{{font-size:26px;margin:0 0 4px}}
.sub{{color:var(--mute);margin:0 0 18px;font-size:14px}}
.changes{{background:var(--card);border:1px solid var(--line);border-left:4px solid var(--accent);border-radius:12px;padding:12px 16px;margin:0 0 20px;font-size:14px}}
.changes b{{display:block;margin-bottom:4px}}
.changes ul{{margin:0;padding-left:18px}}
input[name=tab]{{position:absolute;opacity:0;pointer-events:none}}
.tabs{{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px}}
.tabs label{{cursor:pointer;padding:8px 14px;border-radius:999px;border:1px solid var(--line);background:var(--card);font-weight:600;font-size:14px}}
.tabs label span{{color:var(--mute);font-weight:400;margin-left:4px}}
.panel{{display:none}}
#t1:checked~.tabs label[for=t1],#t2:checked~.tabs label[for=t2],#t3:checked~.tabs label[for=t3]{{background:var(--accent);border-color:var(--accent);color:#fff}}
#t1:checked~.tabs label[for=t1] span,#t2:checked~.tabs label[for=t2] span,#t3:checked~.tabs label[for=t3] span{{color:#fff}}
#t1:checked~#p1,#t2:checked~#p2,#t3:checked~#p3{{display:block}}
.note{{color:var(--mute);font-size:13px;margin:0 0 10px}}
.wrap{{overflow-x:auto;background:var(--card);border:1px solid var(--line);border-radius:12px}}
table{{border-collapse:collapse;width:100%;min-width:980px}}
th{{text-align:left;font-size:12px;text-transform:uppercase;letter-spacing:.03em;color:var(--mute);padding:10px;border-bottom:1px solid var(--line);background:var(--accent-soft);position:sticky;top:0}}
td{{padding:10px;border-bottom:1px solid var(--line);vertical-align:top;font-size:14px}}
tr:last-child td{{border-bottom:0}}
td small{{display:block;color:var(--mute);font-size:12px}}
td.n{{color:var(--mute);width:28px}}
td.app{{width:30%}} td.app span{{display:block;color:var(--mute);font-size:13px;margin-top:2px}}
td.app .lk{{margin-top:4px}} td.app a{{color:var(--accent);text-decoration:none}} td.app a:hover{{text-decoration:underline}}
td.num{{white-space:nowrap;font-variant-numeric:tabular-nums}}
td.why{{font-size:13px;width:28%}}
.z{{display:inline-block;width:8px;height:8px;border-radius:50%;margin-right:6px;vertical-align:middle}}
.z.lo{{background:var(--lo)}} .z.mid{{background:var(--mid)}} .z.hi{{background:var(--hi)}} .z.na{{background:var(--na)}}
.t{{font-style:normal;font-size:12px;padding:2px 8px;border-radius:999px;white-space:nowrap;background:var(--accent-soft)}}
.t-растёт{{background:#d9f0dd;color:#2f6b3a}} .t-держится{{background:#fbefcf;color:#7a5a12}} .t-падает{{background:#f8d6da;color:#8a2333}}
.t-новый,.t-вернулся{{background:#dbe8f6;color:#2c4f7a}}
</style></head>
<body><main>
<h1>Шорт-лист приложений по FB-выгрузкам</h1>
<p class="sub">Метрики — выгрузка 10.10 (web2web giants #9), спенд за 28 дней. Тренд — октябрь к августу; ранние выгрузки (май–август) — только для тренда. Обновлено 28.10.</p>
<div class="changes"><b>Что изменилось 28.10</b><ul>
<li>Сводная сверена построчно со всеми 6 выгрузками — цифры перенесены верно; исправлены расчёты поверх: тренды, август у части приложений, скоры, порядок.</li>
<li>Шорт-лист — только приложения из выгрузки 10.10 с веб-воронкой и оплатами: {len(S)} приложений. Убраны Mosso, Aligned и Chillio (Aligned и Chillio — в «Под вопросом»).</li>
<li>У Chillio в выгрузке 10.10 оплат нет (ROAS 0) — CPA сравнивать нельзя; оплаты видны только в мае.</li>
</ul></div>
<input type="radio" name="tab" id="t1" checked><input type="radio" name="tab" id="t2"><input type="radio" name="tab" id="t3">
<div class="tabs"><label for="t1">Шорт-лист<span>{len(S)}</span></label><label for="t2">Под вопросом<span>{len(Q)}</span></label><label for="t3">Топ-кандидаты<span>{len(T)}</span></label></div>
<section class="panel" id="p1"><p class="note">Кандидаты без оговорок, по скору: 30% спенд, 25% доля tier-1, 25% дешевизна CPA, 20% тренд. Точка у CPA — треть выгрузки 10.10: зелёная — дёшево, жёлтая — средне, красная — дорого.</p>{table(S, 'Почему')}</section>
<section class="panel" id="p2"><p class="note">Заметный спенд в выгрузке 10.10, но есть сомнения по нише, CPA или качеству данных.</p>{table(Q, 'Почему под вопросом')}</section>
<section class="panel" id="p3"><p class="note">Все кандидаты основного блока (без health): спенд в 10.10 ≥ 250k, тренд растёт/держится/новый, CPA в зелёной или жёлтой зоне. По скору.</p>{table(T, 'Почему')}</section>
</main></body></html>'''
open(f'{ROOT}/reports/fb-shortlist-28-10.html', 'w', encoding='utf-8').write(page)
print('ok', len(S), len(Q), len(T))
