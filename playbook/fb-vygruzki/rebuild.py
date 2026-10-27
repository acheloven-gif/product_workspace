import openpyxl,collections,math,re,copy
exec(open('agg.py').read())
OUT='/home/user/product-workspace/reports/fb-kandidaty-itog-27-10.xlsx'
MONTHS=['Май','Июнь','Июль','Август','Октябрь']
blk={r[0]:r[1] for r in rows}
apps=sorted(blk)
def main_scheme(a):
    return 'веб' if any(summarize(a,m,'веб') for m in MONTHS) else 'стор'
M={}
for a in apps:
    s=main_scheme(a); d={'scheme':s}
    for m in MONTHS:
        x=summarize(a,m,s); d[m]=x
    d['other']={m:summarize(a,m,'стор' if s=='веб' else 'веб') for m in MONTHS}
    M[a]=d
def sp(a,m): x=M[a][m]; return x['spend'] if x else None
def cpa(a,m): x=M[a][m]; return x['cpa'] if x else None
oct_vals=[cpa(a,'Октябрь') for a in apps if cpa(a,'Октябрь')]
aug_vals=[cpa(a,'Август') for a in apps if cpa(a,'Август')]
def pct(v,vals):
    if v is None: return None
    return (sum(1 for x in vals if x<v)+1)/len(vals)
def ratio(a,b): return (a/b-1) if a and b else None
def trend(r,new):
    if new: return 'новый'
    if r is None: return None
    return 'растёт' if r>0.2 else ('падает' if r<-0.2 else 'держится')
for a in apps:
    d=M[a]
    d['jun_aug']= trend(ratio(sp(a,'Август'), sp(a,'Июль') or sp(a,'Июнь')), not (sp(a,'Июль') or sp(a,'Июнь')) and sp(a,'Август')) if sp(a,'Август') else None
    d['aug_oct']= trend(ratio(sp(a,'Октябрь'),sp(a,'Август')), not sp(a,'Август')) if sp(a,'Октябрь') else 'пропал'
    d['pct_oct']=pct(cpa(a,'Октябрь'),oct_vals); d['pct_aug']=pct(cpa(a,'Август'),aug_vals)
TW={'растёт':1,'держится':0.7,'новый':0.5}
def score(a,t1,tr=None):
    s=sp(a,'Октябрь') if a in MIXED else (OCT_ROW.get(a) or sp(a,'Октябрь'))
    if not s or t1 is None: return None
    p=M[a]['pct_oct'] if M[a]['pct_oct'] is not None else 1
    return round(30*min(max((math.log10(s)-5)/2,0),1)+25*t1+25*(1-p)+20*TW.get(tr or M[a]['aug_oct'],0),1)
def fmtM(v): return '—' if not v else f'{v/1e6:.2f}M'
def zone(p): return None if p is None else ('в нижней трети (дёшево)' if p<=1/3 else ('в средней трети' if p<=2/3 else 'в верхней трети (дорого)'))
def mixed_note(a):
    o=M[a]['other']; parts=[f"{m[:3].lower()} {o[m]['spend']/1e6:.2f}M, CPA {o[m]['cpa']:.1f}" for m in MONTHS if o[m] and o[m]['cpa']]
    if not parts: return ''
    other='стор' if M[a]['scheme']=='веб' else 'веб'
    return f" Отдельно схема «{'прямо в стор' if other=='стор' else 'веб-воронка'}» (в расчёт не входит, конверсия там — установка): " + '; '.join(parts)+'.'
wb=openpyxl.load_workbook(SRC)
changes=[]
DELTA={}
OCT_ROW={}
for _n in ['Топ-кандидаты','Находки (gems)']:
    _w=wb[_n]
    for _r in range(4,_w.max_row+1):
        if _w.cell(_r,2).value: OCT_ROW[_w.cell(_r,2).value]=_w.cell(_r,13).value
MIXED={a for a in M if any(M[a]['other'][m] for m in MONTHS)}
KEEP_FOR_PURE=('Спенд','Июнь','Июль','Август','Октябрь','Май','Тренд','Август к июлю','Октябрь к августу')
def setv(ws,r,c,v,a,label):
    old=ws.cell(r,c).value
    if isinstance(old,str) and old.startswith('='): return
    if a not in MIXED and label.startswith(KEEP_FOR_PURE): return
    if (isinstance(old,(int,float)) and isinstance(v,(int,float)) and abs((old or 0)-v)>max(1e-6,abs(v)*0.005)) or (not isinstance(v,(int,float)) and old!=v):
        changes.append((ws.title,a,label,old,v))
    ws.cell(r,c).value=v
SCH={'веб':'веб-воронка (квиз → оплата на сайте)','стор':'прямо в стор'}
# ---- Top & Gems (same column layout)
for name in ['Топ-кандидаты','Находки (gems)']:
    ws=wb[name]
    hdr=[c.value for c in ws[3]]
    col={h:i+1 for i,h in enumerate(hdr) if h}
    for r in range(4,ws.max_row+1):
        a=ws.cell(r,2).value
        if a not in M: continue
        d=M[a]
        OLD=(ws.cell(r,col['Спенд 28д, Октябрь (10.10)']).value,ws.cell(r,col['CPA-перцентиль окт (0 = дешевле всех)']).value,ws.cell(r,col['Тренд авг→окт']).value,ws.cell(r,col['Скор (0–100)']).value)
        for m,h in [('Июнь','Спенд 28д, Июнь'),('Июль','Спенд 28д, Июль'),('Август','Спенд 28д, Август'),('Октябрь','Спенд 28д, Октябрь (10.10)')]:
            setv(ws,r,col[h],sp(a,m),a,h)
        setv(ws,r,col['Август к июлю'],ratio(sp(a,'Август'),sp(a,'Июль')),a,'Август к июлю')
        setv(ws,r,col['Октябрь к августу'],ratio(sp(a,'Октябрь'),sp(a,'Август')),a,'Октябрь к августу')
        setv(ws,r,col['CPA окт'],round(cpa(a,'Октябрь'),1) if cpa(a,'Октябрь') else None,a,'CPA окт')
        setv(ws,r,col['CPA-перцентиль окт (0 = дешевле всех)'],d['pct_oct'],a,'CPA-перцентиль')
        setv(ws,r,col['Тренд июнь–авг'],d['jun_aug'],a,'Тренд июнь–авг')
        setv(ws,r,col['Тренд авг→окт'],d['aug_oct'],a,'Тренд авг→окт')
        if isinstance(OLD[3],(int,float)) and OLD[0]:
            def part(spv,p,tr): return 30*min(max((math.log10(spv)-5)/2,0),1)+25*(1-(p if p is not None else 1))+20*TW.get(tr,0)
            newp=part(ws.cell(r,col['Спенд 28д, Октябрь (10.10)']).value or OLD[0],ws.cell(r,col['CPA-перцентиль окт (0 = дешевле всех)']).value,ws.cell(r,col['Тренд авг→окт']).value)
            oldp=part(OLD[0],OLD[1],OLD[2])
            DELTA[a]=newp-oldp
            setv(ws,r,col['Скор (0–100)'],round(OLD[3]+newp-oldp,1),a,'Скор')
        if any(d['other'][m] for m in MONTHS):
            ws.cell(r,col['Куда ведёт реклама']).value=SCH[d['scheme']]+' — в расчёте; стор-схема отдельно'
        why=[h for h in hdr if h and h.startswith('Почему')][0]
        t=ws.cell(r,col[why]).value or ''
        if t.startswith('Спенд 28д') and a in MIXED:
            t=f"Спенд 28д ({'веб-воронка' if d['scheme']=='веб' else 'стор'}): июль {fmtM(sp(a,'Июль'))} → авг {fmtM(sp(a,'Август'))} → окт {fmtM(sp(a,'Октябрь'))} ({d['aug_oct']}). CPA в октябре {zone(d['pct_oct'])}."
        t=re.sub(r' Отдельно схема «.*$','',t)+mixed_note(a)
        ws.cell(r,col[why]).value=t
    if name=='Топ-кандидаты':
        data=[[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(4,ws.max_row+1) if ws.cell(r,2).value in M]
        links=[[ws.cell(r,c).hyperlink for c in range(1,ws.max_column+1)] for r in range(4,ws.max_row+1) if ws.cell(r,2).value in M]
        order=sorted(range(len(data)),key=lambda i:-(data[i][5] or 0))
        for k,i in enumerate(order):
            for c,v in enumerate(data[i],1):
                cell=ws.cell(4+k,c); cell.value=v; cell.hyperlink=None
                if links[i][c-1] is not None and links[i][c-1].target: cell.hyperlink=links[i][c-1].target
            ws.cell(4+k,1).value=k+1
# ---- Health
ws=wb['Health и фитнес']; hdr=[c.value for c in ws[1]]; col={h:i+1 for i,h in enumerate(hdr) if h}
for r in range(2,ws.max_row+1):
    a=ws.cell(r,1).value
    if a not in M: continue
    d=M[a]
    for m in MONTHS: setv(ws,r,col[f'Спенд 28д, {m}'],sp(a,m),a,m)
    setv(ws,r,col['Август к июлю'],ratio(sp(a,'Август'),sp(a,'Июль')),a,'Август к июлю')
    setv(ws,r,col['Октябрь к августу'],ratio(sp(a,'Октябрь'),sp(a,'Август')),a,'Октябрь к августу')
    setv(ws,r,col['CPA авг'],round(cpa(a,'Август'),1) if cpa(a,'Август') else None,a,'CPA авг')
    setv(ws,r,col['CPA-перцентиль авг'],d['pct_aug'],a,'pct авг')
    setv(ws,r,col['CPA окт'],round(cpa(a,'Октябрь'),1) if cpa(a,'Октябрь') else None,a,'CPA окт')
    setv(ws,r,col['CPA-перцентиль окт'],d['pct_oct'],a,'pct окт')
    setv(ws,r,col['Тренд июнь–авг'],d['jun_aug'],a,'Тренд июнь–авг')
    setv(ws,r,col['Тренд авг→окт'],d['aug_oct'],a,'Тренд авг→окт')
    if any(d['other'][m] for m in MONTHS):
        ws.cell(r,col['Куда ведёт реклама']).value=SCH[d['scheme']]+' — в расчёте; стор-схема отдельно (лист «Веб и стор отдельно»)'
# ---- Shortlist / question: contents + card lines
tier={}
for name in ['Топ-кандидаты','Находки (gems)']:
    w=wb[name]
    for r in range(4,w.max_row+1):
        if w.cell(r,2).value in M: tier[w.cell(r,2).value]=w.cell(r,9).value
def resolve(nm):
    if nm in M: return nm
    for a in M:
        if nm.startswith(a) or a.startswith(nm.split(' (')[0]): return a
for name in ['Шорт-лист','Под вопросом']:
    ws=wb[name]
    r=6
    while ws.cell(r,2).value:
        a=resolve(ws.cell(r,2).value)
        if a:
            setv(ws,r,4,sp(a,'Октябрь'),a,'Спенд 28д окт')
            setv(ws,r,5,round(cpa(a,'Октябрь'),1) if cpa(a,'Октябрь') else None,a,'CPA окт')
            setv(ws,r,7,M[a]['aug_oct'],a,'Тренд')
            t1=tier.get(a)
            if t1 is None:
                # tier share from card line
                pass
            if a in DELTA and isinstance(ws.cell(r,3).value,(int,float)): setv(ws,r,3,round(ws.cell(r,3).value+DELTA[a]),a,'Скор')
        r+=1
    for rr in range(r,ws.max_row+1):
        v=ws.cell(rr,1).value
        if not isinstance(v,str): continue
        if ' · срез 10.10 · ' in v and 'спенд 28д по всем' in v:
            a=resolve(ws.cell(rr-1,1).value or '')
            if a and any(M[a]['other'][m] for m in MONTHS):
                o=M[a]['other']['Октябрь']
                v=re.sub(r'спенд 28д по всем — [\d  ]+\$', f"спенд 28д, веб-воронка — {sp(a,'Октябрь'):,.0f} $".replace(',',' ')+(f"; прямо в стор — {o['spend']:,.0f} $ (отдельная схема, в расчёт не входит)".replace(',',' ') if o else ''),v)
                ws.cell(rr,1).value=v
        if v.startswith('Спенд приложения за 28 дней по срезам'):
            # find app: scan up to card title
            k=rr
            while k>r and not (ws.cell(k,1).value and ws.cell(k+1,1).value and ' · срез ' in str(ws.cell(k+1,1).value)): k-=1
            a=resolve(ws.cell(k,1).value or '')
            if a:
                d=M[a]
                main=' → '.join(f"{m}: {d[m]['spend']/1e6:.2f} млн $ (CPA {d[m]['cpa']:.0f})" for m in MONTHS if d[m] and d[m]['cpa'])
                txt=f"Спенд приложения за 28 дней по срезам ({'веб-воронка' if d['scheme']=='веб' else 'стор'}, CPA = спенд / количество конверсий): {main}"
                if any(d['other'][m] for m in MONTHS):
                    txt+='. Отдельно «прямо в стор» (конверсия = установка): '+' → '.join(f"{m}: {d['other'][m]['spend']/1e6:.2f} млн $ (CPA {d['other'][m]['cpa']:.1f})" for m in MONTHS if d['other'][m] and d['other'][m]['cpa'])
                ws.cell(rr,1).value=txt
# ---- Сводная: scheme column next to CPA
ws=wb['Сводная (все строки)']
links={}
for row in ws.iter_rows():
    for c in row:
        if c.hyperlink and c.hyperlink.target: links[(c.row,c.column)]=c.hyperlink.target
ws.insert_cols(18)
for (rr,cc),t in links.items():
    ws.cell(rr,cc).hyperlink=None
for (rr,cc),t in links.items():
    nc=cc+1 if cc>=18 else cc
    ws.cell(rr,nc).hyperlink=t
ws.cell(1,18).value='куда ведёт реклама (схема)'
ws.cell(1,18).font=copy.copy(ws.cell(1,17).font); ws.cell(1,18).fill=copy.copy(ws.cell(1,17).fill)
for rr in range(2,ws.max_row+1):
    ws.cell(rr,18).value=scheme(ws.cell(rr,49).value)
ws.column_dimensions['R'].width=16
# ---- New sheet: web vs store
ns=wb.create_sheet('Веб и стор отдельно',1)
ns.append(['Приложения, у которых в выгрузках одновременно две схемы: веб-воронка (конверсия = оплата/лид на сайте) и реклама прямо в стор (конверсия = установка). Схемы не складываем: в остальных листах — только основная схема (веб-воронка). CPA = спенд за 28 дней / «количество конверсий» из выгрузки.'])
ns.append([])
ns.append(['Приложение','Схема']+[f'Спенд 28д, {m}' for m in MONTHS]+[f'CPA, {m}' for m in MONTHS]+['Октябрь к августу'])
for a in apps:
    d=M[a]
    if not any(d['other'][m] for m in MONTHS): continue
    for s,src in [(d['scheme'],{m:d[m] for m in MONTHS}),('стор' if d['scheme']=='веб' else 'веб',d['other'])]:
        ns.append([a,SCH[s]]+[src[m]['spend'] if src[m] else None for m in MONTHS]+[round(src[m]['cpa'],2) if src[m] and src[m]['cpa'] else None for m in MONTHS]+[ratio(src['Октябрь']['spend'] if src['Октябрь'] else None, src['Август']['spend'] if src['Август'] else None)])
for c,wdt in zip('ABCDEFGHIJKLM',[28,34,14,14,14,14,14,10,10,10,10,10,14]): ns.column_dimensions[c].width=wdt
for row in ns.iter_rows(min_row=4):
    for c in row[2:7]: c.number_format='#,##0'
    row[12].number_format='0%'
ns['A1'].alignment=openpyxl.styles.Alignment(wrap_text=True); ns.merge_cells('A1:M1'); ns.row_dimensions[1].height=45
for c in ns[3]: c.font=openpyxl.styles.Font(bold=True)
# ---- Methodology sheet
ms=wb.create_sheet('Методика (27.10)',0)
for line in ['Методика расчёта (обновлено 27.10 по комментарию Кати от 21.10)',
 '1. Строки с лендом в сторе (apps.apple.com / itunes / play.google) не суммируются со строками с веб-лендом: «конверсия» в них — установка, а не оплата. Во всех сводных листах считается основная схема — веб-воронка; стор-схема — на листе «Веб и стор отдельно».',
 '2. В «Сводной (все строки)» рядом с CPA добавлена колонка «куда ведёт реклама (схема)»: веб / стор.',
 '3. CPA по приложению = спенд за 28 дней / «количество конверсий» из выгрузки (не восстанавливается из округлённого CPA).',
 '4. Каждая строка исходника остаётся отдельной строкой в «Сводной»; агрегаты — только внутри одной схемы. Если один пиксель есть в двух августовских выгрузках (27.08 и 28.08), берётся одна (с бóльшим спендом).',
 '5. Пересчитаны: спенд по месяцам, CPA, тренды, CPA-перцентиль и скор в листах «Топ-кандидаты», «Находки», «Health и фитнес», таблицах содержания и строках спенда в карточках «Шорт-лист» и «Под вопросом». Порядок карточек в «Шорт-листе» и «Под вопросом» не менялся — актуальный скор в таблице содержания. Лист «Вертикали» — суммарный спенд всех схем, без изменений.',
 '6. Пример memoryOS: CPA веб-воронки в октябре — $65 (было 8.6 — смешаны оплаты и установки); спенд веб-воронки август 246 357 → октябрь 347 503, рост +41%.',
 '','Что изменилось (было → стало):']:
    ms.append([line])
ms.column_dimensions['A'].width=40
for c in range(2,6): ms.column_dimensions['ABCDE'[c-1]].width=22
ms.append(['Лист','Приложение','Поле','Было','Стало'])
seen=set()
for t,a,l,o,n in changes:
    if (t,a,l) in seen: continue
    seen.add((t,a,l)); ms.append([t,a,l,o,n])
for r in range(2,8): ms.cell(r,1).alignment=openpyxl.styles.Alignment(wrap_text=True)
wb.save(OUT)
print(len(changes),'changes'); 
import collections as C
print(C.Counter((t) for t,*_ in changes))
for t,a,l,o,n in changes:
    if l in ('CPA окт','Скор','Спенд 28д окт') and t in ('Шорт-лист','Под вопросом','Топ-кандидаты'): print(t,a,l,o,'→',n)
